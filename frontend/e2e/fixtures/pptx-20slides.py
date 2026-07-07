"""为端到端测试生成 20 张固定内容的 PPTX 课件素材。

供 front-viewer.spec.ts 的 F-050 / F-053 用例使用:
  - 初始页 1 / 20
  - 末页点 `p-viewer-ppt-next` 按钮 disabled + 文案变 `已是最后一张`
  - 全屏切换 + 键盘 ArrowLeft / ArrowRight 切换

运行:
  pip install python-pptx
  python fixtures/pptx-20slides.py

输出:
  - fixtures/courseware-with-pptx-20slides.pptx  (20 张,标题 Slide N)
  - fixtures/courseware-with-pptx-1slide.pptx    (1 张,验证首末页 disabled)
  - fixtures/corrupt.pptx                        (损坏文件,验证失败降级)
"""
from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:  # pragma: no cover
    print(
        "缺少 python-pptx 依赖, 请先 `pip install python-pptx`",
        file=sys.stderr,
    )
    raise


HERE = Path(__file__).resolve().parent


def build_20_slides_pptx(out: Path) -> None:
    prs = Presentation()
    # 使用内置空白版式
    blank = prs.slide_layouts[6]
    for i in range(1, 21):
        slide = prs.slides.add_slide(blank)
        # 写标题
        tx = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
        tf = tx.text_frame
        tf.text = f"Slide {i}"
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(36)
                r.font.bold = True
        # 写正文
        body = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(4))
        btf = body.text_frame
        btf.text = (
            f"这是第 {i} / 20 张幻灯片\n"
            f"用于端到端测试 front-viewer.spec.ts F-050 / F-053\n"
            f"验证翻页、键盘 ←/→ 切换、首末页 disabled 与全屏切换。"
        )
        for p in btf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(18)
    prs.save(out)
    print(f"[ok] 生成 20 张 PPTX: {out}  size={out.stat().st_size} bytes")


def build_1_slide_pptx(out: Path) -> None:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    tx = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tf = tx.text_frame
    tf.text = "Single Slide"
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(36)
            r.font.bold = True
    prs.save(out)
    print(f"[ok] 生成 1 张 PPTX: {out}  size={out.stat().st_size} bytes")


def build_corrupt_pptx(out: Path) -> None:
    """构造一个会被 LibreOffice 拒绝的损坏 PPTX, 用于测试转换失败降级。

    PPTX 实际是 zip 容器, 这里写入非 zip 字节但扩展名为 .pptx。
    """
    out.write_bytes(b"NOT_A_VALID_PPTX_BUT_HAS_PPTX_EXT\n" * 8)
    print(f"[ok] 生成 corrupt PPTX: {out}  size={out.stat().st_size} bytes")


def main() -> int:
    HERE.mkdir(parents=True, exist_ok=True)
    build_20_slides_pptx(HERE / "courseware-with-pptx-20slides.pptx")
    build_1_slide_pptx(HERE / "courseware-with-pptx-1slide.pptx")
    build_corrupt_pptx(HERE / "corrupt.pptx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
